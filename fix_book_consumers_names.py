"""
Regenerates 06_book_consumers_156k.xlsx with intelligent first-name extraction
from usernames.

Pattern handling:
  nicholas.c.devito  → first segment before dot   → Nicholas
  taylor.hilliard    → first segment before dot   → Taylor
  mi.zeena           → first segment before dot   → Mi
  fergs_plumbing     → first segment before _     → Fergs
  poonam-saini       → first segment before -     → Poonam
  jadegarciamakeup   → longest known-name prefix  → Jade
  lucaslokesh        → longest known-name prefix  → Lucas
  jmatz              → no match → whole username  → Jmatz (fallback)
"""
import csv, re, os
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from names_dataset import NameDataset

EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
TYPO_MAP = {
    "gmal.com":"gmail.com","gmial.com":"gmail.com","gamil.com":"gmail.com",
    "gmail.co":"gmail.com","gnail.com":"gmail.com","gmil.com":"gmail.com",
    "gmaill.com":"gmail.com","gmai.com":"gmail.com",
    "yaho.com":"yahoo.com","yahooo.com":"yahoo.com","yhoo.com":"yahoo.com",
    "hotmal.com":"hotmail.com","hotmial.com":"hotmail.com","hotmai.com":"hotmail.com",
    "outlok.com":"outlook.com","outllok.com":"outlook.com","outook.com":"outlook.com",
    "icoud.com":"icloud.com","iclod.com":"icloud.com",
}

def normalize(email):
    e = email.strip().lower()
    if "@" not in e: return e
    local, _, domain = e.rpartition("@")
    return f"{local}@{TYPO_MAP.get(domain, domain)}"

def fix_email(email):
    if "@" not in email: return email
    local, _, domain = email.rpartition("@")
    return f"{local}@{TYPO_MAP.get(domain.lower(), domain.lower())}"

BASE = "/home/user/alx-pre_course"
OUT  = f"{BASE}/final_lists"

# ── Build first-name lookup set (single words only, lowercase) ────────────────
print("Loading name dictionary...")
nd = NameDataset()
FIRST_NAMES = {k.lower() for k in nd.first_names if " " not in k}
print(f"  {len(FIRST_NAMES):,} single-word first names loaded")

def extract_firstname(username: str) -> str:
    """Intelligently extract a first name from a username."""
    u = username.strip()
    if not u:
        return ""
    ul = u.lower()

    # Step 1: separator present → first segment
    for sep in [".", "_", "-"]:
        if sep in ul:
            seg = ul.split(sep)[0].strip()
            return seg.capitalize() if seg else u.capitalize()

    # Step 2: concatenated → find longest known-first-name prefix
    # scan lengths 2..12 and keep the longest hit
    best = None
    for length in range(2, min(len(ul) + 1, 13)):
        prefix = ul[:length]
        if prefix in FIRST_NAMES:
            best = prefix            # keep extending to find longest match

    if best:
        return best.capitalize()

    # Step 3: fallback — capitalise the whole username
    return u.capitalize()

# ── Load previously established global_seen (same priority order) ─────────────
# We need to know which emails in this file are globally unique.
# Re-parse all lists to rebuild global_seen.
SOURCES = [
    ("LinkedIn Leads",      f"{BASE}/leads_raw.csv",               "named"),
    ("Educators",           f"{BASE}/leads2_raw.csv",              "split"),
    ("US Bookstores",       f"{BASE}/leads3_raw.csv",              "numeric_book"),
    ("CBS MarketWatch",     f"{BASE}/leads4_raw.csv",              "numeric_cbs"),
    ("Leads",               f"{BASE}/leads5_raw.csv",              "split"),
    ("Book Consumers 156K", f"{BASE}/leads6_full.csv",             "username"),
]

def parse_email_only(filepath, fmt):
    norms = []
    with open(filepath, encoding="utf-8-sig", newline="") as f:
        for raw in csv.DictReader(f):
            try:
                if fmt == "named":      e = raw.get("Email","") or ""
                elif fmt == "split":    e = raw.get("Email","") or ""
                elif fmt == "numeric_book": e = raw.get("11","") or ""
                elif fmt == "numeric_cbs":  e = raw.get("0","") or ""
                elif fmt == "username": e = raw.get("Email","") or ""
                else: continue
                e = e.strip()
                if EMAIL_RE.match(e): norms.append(normalize(fix_email(e)))
            except Exception: pass
    return norms

print("Building global dedup map...")
global_seen = {}
for label, fp, fmt in SOURCES:
    for ne in parse_email_only(fp, fmt):
        if ne not in global_seen:
            global_seen[ne] = label
print(f"  {len(global_seen):,} unique emails indexed")

# ── Parse Book Consumers with smart name extraction ───────────────────────────
print("Parsing Book Consumers 156K with smart name extraction...")
rows = []
with open(f"{BASE}/leads6_full.csv", encoding="utf-8-sig", newline="") as f:
    for raw in csv.DictReader(f):
        try:
            username = (raw.get("Username","") or "").strip()
            email    = (raw.get("Email","") or "").strip()
            if not email or not EMAIL_RE.match(email): continue
            email = fix_email(email)
            rows.append({
                "email":            email,
                "first_name":       extract_firstname(username),
                "last_name":        "",
                "company_name":     "",
                "phone_number":     "",
                "website":          "",
                "linkedin_profile": "",
                "location":        "",
                "norm":             normalize(email),
            })
        except Exception: pass

print(f"  {len(rows):,} rows parsed")

# Keep only globally unique contacts owned by this list
seen = set()
unique_rows = []
for row in rows:
    ne = row["norm"]
    if global_seen.get(ne) != "Book Consumers 156K": continue
    if ne in seen: continue
    seen.add(ne)
    unique_rows.append(row)

print(f"  {len(unique_rows):,} globally unique contacts")

# Quick sample of name extraction results
import random
sample = random.sample(unique_rows, min(20, len(unique_rows)))
print("\n  Sample name extractions:")
for r in sample:
    print(f"    {r['first_name']:<20s}  ← email: {r['email'][:40]}")

# ── Write Excel ────────────────────────────────────────────────────────────────
TEMPLATE_COLS = [
    "email", "first_name", "last_name", "company_name",
    "phone_number", "website", "linkedin_profile", "location",
]
COL_WIDTHS = {
    "email":36, "first_name":20, "last_name":20, "company_name":28,
    "phone_number":18, "website":30, "linkedin_profile":36, "location":24,
}

FH   = PatternFill("solid", fgColor="2E5090")
FALT = PatternFill("solid", fgColor="EEF3FB")
WB   = Font(bold=True, color="FFFFFF")
THIN = Border(left=Side(style="thin"), right=Side(style="thin"),
              top=Side(style="thin"),  bottom=Side(style="thin"))
CTR  = Alignment(horizontal="center", vertical="center")
WRP  = Alignment(vertical="center", wrap_text=False)

print("\nWriting 06_book_consumers_156k.xlsx ...")
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Contacts"

for i, col in enumerate(TEMPLATE_COLS, 1):
    cell = ws.cell(1, i, col)
    cell.fill = FH; cell.font = WB; cell.alignment = CTR; cell.border = THIN
    ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS.get(col, 20)
ws.row_dimensions[1].height = 26
ws.freeze_panes = "A2"

for ri, row in enumerate(unique_rows, 2):
    for ci, col in enumerate(TEMPLATE_COLS, 1):
        cell = ws.cell(ri, ci, row.get(col, ""))
        cell.alignment = WRP; cell.border = THIN
        if ri % 2 == 0:
            cell.fill = FALT
    ws.row_dimensions[ri].height = 15

wb.save(f"{OUT}/06_book_consumers_156k.xlsx")
print(f"  Saved: {OUT}/06_book_consumers_156k.xlsx")
print(f"\nDone — {len(unique_rows):,} contacts written.")
