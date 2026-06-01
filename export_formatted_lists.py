"""
Exports clean, globally-deduplicated contact lists in standardised template:
  email | first_name | last_name | company_name | phone_number |
  website | linkedin_profile | location

- Single sheet per file (no tabs)
- Only globally unique contacts (first-list-wins dedup)
- Columns skipped if all empty in that list
- Batch 7 gets a single sheet with a "segment" column
- Master gets a "source_list" column
"""
import csv, re, os
from collections import defaultdict
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

EMAIL_RE = re.compile(r'^[^@\s]+@[^@\s]+\.[^@\s]+$')
TYPO_MAP = {
    "gmal.com":"gmail.com","gmial.com":"gmail.com","gamil.com":"gmail.com",
    "gmail.co":"gmail.com","gnail.com":"gmail.com","gmil.com":"gmail.com",
    "gmaill.com":"gmail.com","gmai.com":"gmail.com",
    "yaho.com":"yahoo.com","yahooo.com":"yahoo.com","yhoo.com":"yahoo.com",
    "hotmal.com":"hotmail.com","hotmial.com":"hotmail.com","hotmai.com":"hotmail.com",
    "outlok.com":"outlook.com","outllok.com":"outlook.com","outook.com":"outlook.com",
    "icoud.com":"icloud.com","iclod.com":"icloud.com",
}

def normalize(email):
    e = email.strip().lower()
    if "@" not in e: return e
    local, _, domain = e.rpartition("@")
    return f"{local}@{TYPO_MAP.get(domain, domain)}"

def fix_email(email):
    if "@" not in email: return email
    local, _, domain = email.rpartition("@")
    return f"{local}@{TYPO_MAP.get(domain.lower(), domain.lower())}"

def split_name(full):
    parts = full.strip().split(None, 1)
    return (parts[0] if parts else ""), (parts[1] if len(parts) > 1 else "")

BASE = "/home/user/alx-pre_course"
OUT  = f"{BASE}/final_lists"
os.makedirs(OUT, exist_ok=True)

SOURCES = [
    ("LinkedIn Leads",      f"{BASE}/leads_raw.csv",               "named",        "01_linkedin_leads.xlsx"),
    ("Educators",           f"{BASE}/leads2_raw.csv",              "split",        "02_educators.xlsx"),
    ("US Bookstores",       f"{BASE}/leads3_raw.csv",              "numeric_book", "03_us_bookstores.xlsx"),
    ("CBS MarketWatch",     f"{BASE}/leads4_raw.csv",              "numeric_cbs",  "04_cbs_marketwatch.xlsx"),
    ("Leads",               f"{BASE}/leads5_raw.csv",              "split",        "05_leads.xlsx"),
    ("Book Consumers 156K", f"{BASE}/leads6_full.csv",             "username",     "06_book_consumers_156k.xlsx"),
    ("B7 Christians US",    f"{BASE}/batch7/christian_us.csv",     "standard",     None),
    ("B7 Christians UK",    f"{BASE}/batch7/christian_uk.csv",     "standard",     None),
    ("B7 Non-Christian US", f"{BASE}/batch7/non_christian_us.csv", "standard",     None),
    ("B7 Non-Christian UK", f"{BASE}/batch7/non_christian_uk.csv", "standard",     None),
    ("B7 Non-Christian AU", f"{BASE}/batch7/non_christian_au.csv", "standard",     None),
    ("B7 Random US",        f"{BASE}/batch7/random_us.csv",        "standard",     None),
    ("B7 Random UK",        f"{BASE}/batch7/random_uk.csv",        "standard",     None),
    ("B7 Random AU",        f"{BASE}/batch7/random_au.csv",        "standard",     None),
    ("B7 Agencies",         f"{BASE}/batch7/agencies.csv",         "standard",     None),
    ("B7 Bookstores",       f"{BASE}/batch7/bookstores.csv",       "standard",     None),
]

B7_SEG = {
    "B7 Christians US":"Christians US",   "B7 Christians UK":"Christians UK",
    "B7 Non-Christian US":"Non-Christian US","B7 Non-Christian UK":"Non-Christian UK",
    "B7 Non-Christian AU":"Non-Christian AU","B7 Random US":"Random US",
    "B7 Random UK":"Random UK",           "B7 Random AU":"Random AU",
    "B7 Agencies":"Agencies",             "B7 Bookstores":"Bookstores",
}

def parse(filepath, fmt, label):
    rows = []
    with open(filepath, encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for raw in reader:
            try:
                phone = company = website = linkedin = location = ""
                if fmt == "standard":
                    name  = (raw.get("Contact Name","") or "").strip()
                    email = (raw.get("Contact Email","") or "").strip()
                    phone = (raw.get("Contact Phone","") or "").strip()
                    company = (raw.get("Contact Company","") or "").strip()
                    if name in ("Contact Name","First Name"): continue
                    if email and not EMAIL_RE.match(email) and EMAIL_RE.match(phone):
                        name = f"{name} {email}".strip(); email = phone; phone = ""
                    first, last = split_name(name)
                elif fmt == "named":
                    name  = (raw.get("Name","") or "").strip()
                    email = (raw.get("Email","") or "").strip()
                    phone = (raw.get("Phone Number", raw.get("Phone","")) or "").strip()
                    company = (raw.get("Company","") or "").strip()
                    first, last = split_name(name)
                elif fmt == "split":
                    first = (raw.get("First name", raw.get("First Name","")) or "").strip()
                    last  = (raw.get("Last name",  raw.get("Last Name",""))  or "").strip()
                    email = (raw.get("Email","") or "").strip()
                elif fmt == "username":
                    uname = (raw.get("Username","") or "").strip()
                    email = (raw.get("Email","") or "").strip()
                    first, last = split_name(uname)
                elif fmt == "numeric_book":
                    name    = (raw.get("8","") or "").strip()
                    email   = (raw.get("11","") or "").strip()
                    phone   = (raw.get("10","") or "").strip()
                    company = (raw.get("0","") or "").strip()
                    first, last = split_name(name)
                elif fmt == "numeric_cbs":
                    first = (raw.get("1","") or "").strip()
                    last  = (raw.get("2","") or "").strip()
                    email = (raw.get("0","") or "").strip()
                    phone = (raw.get("7","") or "").strip()
                else:
                    continue

                if not email or not EMAIL_RE.match(email): continue
                email = fix_email(email)
                rows.append({
                    "email":            email,
                    "first_name":       first,
                    "last_name":        last,
                    "company_name":     company,
                    "phone_number":     phone,
                    "website":          website,
                    "linkedin_profile": linkedin,
                    "location":         location,
                    "segment":          B7_SEG.get(label, ""),
                    "source_list":      label,
                    "norm":             normalize(email),
                })
            except Exception:
                pass
    return rows

# ── Load all sources ──────────────────────────────────────────────────────────
print("Loading sources...")
all_by_label = {}
for label, fp, fmt, _ in SOURCES:
    rows = parse(fp, fmt, label)
    all_by_label[label] = rows
    print(f"  {label:<25s}: {len(rows):>7,}")

# Global dedup (first list wins)
global_seen = {}
for label, *_ in SOURCES:
    for row in all_by_label[label]:
        ne = row["norm"]
        if ne not in global_seen:
            global_seen[ne] = label

print(f"\nGlobally unique: {len(global_seen):,}")

def unique_rows(label, rows):
    seen = set()
    out = []
    for row in rows:
        ne = row["norm"]
        if global_seen.get(ne) != label: continue
        if ne in seen: continue
        seen.add(ne)
        out.append(row)
    return out

# ── Styling ───────────────────────────────────────────────────────────────────
FH   = PatternFill("solid", fgColor="2E5090")
FALT = PatternFill("solid", fgColor="EEF3FB")
WB   = Font(bold=True, color="FFFFFF")
BLD  = Font(bold=True)
THIN = Border(left=Side(style="thin"), right=Side(style="thin"),
              top=Side(style="thin"),  bottom=Side(style="thin"))
CTR  = Alignment(horizontal="center", vertical="center")
WRP  = Alignment(vertical="center", wrap_text=False)

TEMPLATE_COLS = [
    "email", "first_name", "last_name", "company_name",
    "phone_number", "website", "linkedin_profile", "location",
]
COL_WIDTHS = {
    "email": 36, "first_name": 20, "last_name": 20, "company_name": 28,
    "phone_number": 18, "website": 30, "linkedin_profile": 36, "location": 24,
    "segment": 22, "source_list": 26,
}

def active_cols(rows, extras=None):
    base = (extras or []) + TEMPLATE_COLS
    return [c for c in base if any(row.get(c) for row in rows)]

def write_sheet(ws, rows, cols):
    for i, col in enumerate(cols, 1):
        cell = ws.cell(1, i, col)
        cell.fill = FH; cell.font = WB; cell.alignment = CTR; cell.border = THIN
        ws.column_dimensions[get_column_letter(i)].width = COL_WIDTHS.get(col, 20)
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"
    for ri, row in enumerate(rows, 2):
        for ci, col in enumerate(cols, 1):
            cell = ws.cell(ri, ci, row.get(col, ""))
            cell.alignment = WRP; cell.border = THIN
            if ri % 2 == 0:
                cell.fill = FALT
        ws.row_dimensions[ri].height = 15

# ── Write standalone lists (01–06) ───────────────────────────────────────────
for label, fp, fmt, outfile in SOURCES:
    if not outfile or label.startswith("B7"): continue
    rows = unique_rows(label, all_by_label[label])
    cols = active_cols(rows)
    print(f"\nWriting {outfile}  →  {len(rows):,} unique contacts  |  cols: {cols}")
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Contacts"
    write_sheet(ws, rows, cols)
    wb.save(f"{OUT}/{outfile}")

# ── Write Batch 7 (single sheet, segment column) ─────────────────────────────
b7_rows = []
for label, *_ in SOURCES:
    if label.startswith("B7"):
        b7_rows += unique_rows(label, all_by_label[label])

cols_b7 = active_cols(b7_rows, extras=["segment"])
print(f"\nWriting 07_batch7_segmented.xlsx  →  {len(b7_rows):,} unique contacts  |  cols: {cols_b7}")
wb7 = openpyxl.Workbook()
ws7 = wb7.active; ws7.title = "Contacts"
write_sheet(ws7, b7_rows, cols_b7)
wb7.save(f"{OUT}/07_batch7_segmented.xlsx")

# ── Write Master ──────────────────────────────────────────────────────────────
master_rows = []
seen_m: set = set()
for label, *_ in SOURCES:
    for row in all_by_label[label]:
        ne = row["norm"]
        if global_seen.get(ne) != label: continue
        if ne in seen_m: continue
        seen_m.add(ne)
        master_rows.append(row)

cols_master = active_cols(master_rows, extras=["source_list"])
print(f"\nWriting 00_MASTER_all_unique_contacts.xlsx  →  {len(master_rows):,} contacts  |  cols: {cols_master}")
wbm = openpyxl.Workbook()
wsm = wbm.active; wsm.title = "All Contacts"
write_sheet(wsm, master_rows, cols_master)
wbm.save(f"{OUT}/00_MASTER_all_unique_contacts.xlsx")

print(f"""
{'='*60}
  EXPORT COMPLETE  —  {len(global_seen):,} globally unique contacts
{'='*60}
  Files in: {OUT}/
    00_MASTER_all_unique_contacts.xlsx
    01_linkedin_leads.xlsx
    02_educators.xlsx
    03_us_bookstores.xlsx
    04_cbs_marketwatch.xlsx
    05_leads.xlsx
    06_book_consumers_156k.xlsx
    07_batch7_segmented.xlsx
{'='*60}
""")
